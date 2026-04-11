#!/usr/bin/env python3
from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any, Dict, Iterable, List

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parents[2]
OUTPUTS_ROOT = PROJECT_ROOT / "data" / "outputs"
XLSX_OUT = OUTPUTS_ROOT / "all_experiment_results.xlsx"
CSV_OUT = OUTPUTS_ROOT / "all_experiment_results.csv"
JSON_OUT = OUTPUTS_ROOT / "all_experiment_results.json"

PREFERRED_COLUMNS = [
    "experiment_id",
    "experiment_name",
    "setting_name",
    "sample_id",
    "company",
    "report_year",
    "category",
    "claim_text",
    "claim_page",
    "metric_name",
    "extracted_value",
    "value_unit",
    "context_text",
    "context_page",
    "evidence_source_type",
    "label",
    "final_label",
    "gold_label",
    "llm_label",
    "llm_reason",
    "is_correct",
    "model",
    "temperature",
    "source_sample_file",
    "input_payload",
    "annotator_1",
    "annotator_2",
    "rationale_short",
    "rationale_short_zh",
    "notes",
]


def load_json(path: Path) -> Any:
    with path.open("r", encoding="utf-8") as f:
        return json.load(f)


def normalize_value(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value


def collect_rows() -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []
    for path in sorted(OUTPUTS_ROOT.glob("setting_*/*_predictions.json")):
        payload = load_json(path)
        if not isinstance(payload, dict):
            continue
        results = payload.get("results") or []
        if not isinstance(results, list):
            continue
        for row in results:
            if isinstance(row, dict):
                rows.append(dict(row))
    return rows


def ordered_columns(rows: Iterable[Dict[str, Any]]) -> List[str]:
    seen = set()
    extras: List[str] = []
    for row in rows:
        for key in row.keys():
            if key not in seen and key not in PREFERRED_COLUMNS:
                seen.add(key)
                extras.append(key)
    return [col for col in PREFERRED_COLUMNS if any(col in row for row in rows)] + sorted(extras)


def write_csv(path: Path, rows: List[Dict[str, Any]], columns: List[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=columns)
        writer.writeheader()
        for row in rows:
            writer.writerow({col: normalize_value(row.get(col)) for col in columns})


def write_xlsx(path: Path, rows: List[Dict[str, Any]], columns: List[str]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "results"
    ws.freeze_panes = "A2"

    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append([normalize_value(row.get(col)) for col in columns])

    for idx, column in enumerate(columns, start=1):
        max_len = len(column)
        for cell in ws[get_column_letter(idx)]:
            if cell.value is None:
                continue
            max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(idx)].width = min(max(max_len + 2, 12), 60)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> None:
    rows = collect_rows()
    if not rows:
        raise FileNotFoundError(f"No prediction summary files found under: {OUTPUTS_ROOT}")

    columns = ordered_columns(rows)
    write_csv(CSV_OUT, rows, columns)
    write_xlsx(XLSX_OUT, rows, columns)
    JSON_OUT.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"[DONE] exported {len(rows)} rows -> {XLSX_OUT}")


if __name__ == "__main__":
    main()
